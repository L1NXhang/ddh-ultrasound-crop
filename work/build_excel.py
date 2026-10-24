# -*- coding: utf-8 -*-
import re
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import sys
sys.path.insert(0, "/workspace/work")
from vocab_data import VOCAB

def load_text(path):
    return open(path).read().lower()

t1 = load_text("/workspace/work/p1.txt")
t2 = load_text("/workspace/work/p2.txt")

def count_word(text, word):
    # 短语按整体子串计数；单词按词边界计数（含复数）
    if " " in word or "-" in word:
        return len(re.findall(re.escape(word), text))
    pat = r"\b" + re.escape(word) + r"(s|es)?\b"
    return len(re.findall(pat, text))

rows = []
for word, ipa, pos, cn, level, papers in VOCAB:
    n1 = count_word(t1, word) if "1" in papers else 0
    n2 = count_word(t2, word) if "2" in papers else 0
    rows.append([word, ipa, pos, cn, level, papers, n1, n2, n1 + n2])

rows.sort(key=lambda r: r[0].lower())

# ---------- 样式 ----------
HEADER_FILL = PatternFill("solid", fgColor="2B6CB0")
HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10.5)
LEVEL_COLORS = {
    "四级": "E6F4EA", "六级": "E8F0FE", "考研": "FEF7E0",
    "学术": "F3E8FD", "专业术语": "FDE8E8",
}
LEVEL_FONTS = {
    "四级": "137333", "六级": "1A56DB", "考研": "B7791F",
    "学术": "6B46C1", "专业术语": "C53030",
}
thin = Side(style="thin", color="D9E2EC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

HEADERS = ["序号", "单词", "音标", "词性", "中文释义", "难度级别", "出现论文", "论文一频次", "论文二频次", "总频次"]
WIDTHS = [7, 26, 24, 11, 42, 11, 11, 11, 11, 10]

def write_sheet(ws, title, data, start_no=1):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(name="微软雅黑", size=14, bold=True, color="1A365D")
    c.alignment = CENTER
    ws.row_dimensions[1].height = 28
    for j, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=2, column=j, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[2].height = 22
    for i, row in enumerate(data):
        r = i + 3
        vals = [start_no + i] + row
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = LEFT if j in (2, 4, 5) else CENTER
        lv = row[4]
        ws.cell(row=r, column=6).fill = PatternFill("solid", fgColor=LEVEL_COLORS.get(lv, "FFFFFF"))
        ws.cell(row=r, column=6).font = Font(name="微软雅黑", size=10.5, color=LEVEL_FONTS.get(lv, "000000"), bold=True)
        ws.row_dimensions[r].height = 20
    for j, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:J{len(data)+2}"

wb = Workbook()

# 说明页
info = wb.active
info.title = "使用说明"
info.column_dimensions["A"].width = 100
lines = [
    ("两篇论文精读词汇清单（四级及以上水平）", 16, True, "1A365D"),
    ("", 11, False, "000000"),
    ("■ 词汇来源", 12, True, "2B6CB0"),
    ("论文一：《Developmental retardation of femoral head size and femoral head ossification in mild and severe developmental dysplasia of the hip in infants》（QIMS 2023;13(1):185-195）", 11, False, "000000"),
    ("论文二：《Deep learning-based no-reference quality assessment of anterior segment ultrasound biomicroscopy panoramic images》（QIMS 2025;15(11):11219-11236）", 11, False, "000000"),
    ("", 11, False, "000000"),
    ("■ 工作表结构", 12, True, "2B6CB0"),
    ("1. 全部词汇：两篇论文所有精选词汇，可按「难度级别」「出现论文」筛选", 11, False, "000000"),
    ("2. 论文一词汇：仅出现在论文一的词汇（DDH 股骨头发育研究）", 11, False, "000000"),
    ("3. 论文二词汇：仅出现在论文二的词汇（UBM 图像质量评估）", 11, False, "000000"),
    ("4. 两篇共有词汇：两篇论文都出现的词汇，优先掌握", 11, False, "000000"),
    ("", 11, False, "000000"),
    ("■ 难度级别说明", 12, True, "2B6CB0"),
    ("四级：大学英语四级核心词｜六级：六级大纲词｜考研：考研核心词｜学术：学术通用词（含构词）｜专业术语：医学/影像/深度学习领域术语", 11, False, "000000"),
    ("", 11, False, "000000"),
    ("■ 使用建议", 12, True, "2B6CB0"),
    ("① 先读对应论文的中文导读，了解大意；② 阅读英文原文前，浏览该论文专属词汇表；③ 遇到生词回查音标和释义；④ 频次高的词（如 >5 次）是论文核心词，务必掌握。", 11, False, "000000"),
    ("", 11, False, "000000"),
    (f"共收录词条：{len(rows)} 个", 12, True, "C53030"),
]
for i, (txt, size, bold, color) in enumerate(lines, 1):
    c = info.cell(row=i, column=1, value=txt)
    c.font = Font(name="微软雅黑", size=size, bold=bold, color=color)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    info.row_dimensions[i].height = 24 if size >= 12 else 20

# 全部词汇
ws_all = wb.create_sheet("全部词汇")
write_sheet(ws_all, f"全部词汇（共 {len(rows)} 词）", rows)

# 论文一
p1 = [r for r in rows if r[5] == "1"]
ws1 = wb.create_sheet("论文一词汇")
write_sheet(ws1, f"论文一词汇（DDH 股骨头发育研究，共 {len(p1)} 词）", p1)

# 论文二
p2 = [r for r in rows if r[5] == "2"]
ws2 = wb.create_sheet("论文二词汇")
write_sheet(ws2, f"论文二词汇（UBM 图像质量评估，共 {len(p2)} 词）", p2)

# 共有
both = [r for r in rows if r[5] == "1+2"]
wsb = wb.create_sheet("两篇共有词汇")
write_sheet(wsb, f"两篇共有词汇（共 {len(both)} 词，优先掌握）", both)

out = "/workspace/论文精读单词清单.xlsx"
wb.save(out)
print("saved:", out)
print("total:", len(rows), "| P1:", len(p1), "| P2:", len(p2), "| both:", len(both))
